"""Convert the project Markdown reports to DOCX (headings, paragraphs, lists, tables, images, formula images).

Supported Markdown subset: #..#### headings, paragraphs, '-' and '1.' lists, pipe tables, ![caption](path),
$$ display formulas (pre-rendered by render_formulas.py), ```text blocks, inline `code` and **bold**.
Usage: python md_to_docx.py input.md output.docx
"""
import hashlib
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FORMULAS = Path(__file__).resolve().parent / "build" / "formulas"
CJK, LATIN, MONO = "SimSun", "Times New Roman", "Consolas"
TEXT_WIDTH = Cm(16.0)


def set_font(run, name=LATIN, east=CJK, size=None, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)
    if size:
        run.font.size = size
    if bold is not None:
        run.bold = bold


def inline(paragraph, text, size=None, bold=False):
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("`"):
            set_font(paragraph.add_run(part[1:-1]), MONO, CJK, size or Pt(9.5), bold)
        elif part.startswith("**"):
            set_font(paragraph.add_run(part[2:-2]), size=size, bold=True)
        else:
            set_font(paragraph.add_run(part), size=size, bold=bold)


def styles(document):
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = LATIN, Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.25
    for level, size in ((0, 18), (1, 15), (2, 13), (3, 11.5), (4, 10.5)):
        style = document.styles["Title" if level == 0 else f"Heading {level}"]
        style.font.name, style.font.size, style.font.bold = LATIN, Pt(size), True
        style.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimHei")
    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2.2)


def table(document, rows):
    cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
    cells = [row for row in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in row)]
    width = max(len(row) for row in cells)
    grid = document.add_table(rows=len(cells), cols=width)
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(cells):
        for c in range(width):
            paragraph = grid.cell(r, c).paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            inline(paragraph, row[c] if c < len(row) else "", size=Pt(9), bold=(r == 0))
    document.add_paragraph()


def image(document, path, caption, width=TEXT_WIDTH):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=width)
    if caption:
        cap = document.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        inline(cap, caption, size=Pt(9))


def formula(document, tex):
    from PIL import Image
    for line in [x.strip() for x in tex.strip().splitlines() if x.strip()]:
        path = FORMULAS / f"{hashlib.sha1(line.encode('utf-8')).hexdigest()[:16]}.png"
        if not path.exists():
            raise FileNotFoundError(f"formula image missing (run render_formulas.py): {line}")
        with Image.open(path) as picture:
            natural = Cm(picture.width / 220 * 2.54)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run().add_picture(str(path), width=min(natural, TEXT_WIDTH))


def convert(source, target):
    source = Path(source)
    document = Document()
    styles(document)
    lines = source.read_text(encoding="utf-8-sig").splitlines()
    i = 0
    ordered = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if not re.match(r"^\s*\d+\.\s", line):
            ordered = 0
        if stripped == "$$":
            end = lines.index("$$", i + 1)
            formula(document, "\n".join(lines[i + 1:end]))
            i = end + 1
            continue
        if stripped.startswith("```"):
            end = next(j for j in range(i + 1, len(lines)) if lines[j].strip().startswith("```"))
            for code in lines[i + 1:end]:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.left_indent = Cm(0.8)
                set_font(paragraph.add_run(code), MONO, CJK, Pt(9))
            i = end + 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            document.add_heading(heading.group(2), level=0 if level == 1 else level - 1)
            i += 1
            continue
        picture = re.match(r"^!\[(.*?)\]\(<?(.*?)>?\)$", stripped)
        if picture:
            image(document, (source.parent / picture.group(2)).resolve(), picture.group(1))
            i += 1
            continue
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            table(document, block)
            continue
        bullet = re.match(r"^(\s*)(-|\d+\.)\s+(.*)$", line)
        if bullet:
            if bullet.group(2) == "-":
                style = "List Bullet" + (" 2" if len(bullet.group(1)) >= 2 else "")
                inline(document.add_paragraph(style=style), bullet.group(3))
            else:
                # Keep the number written in the source: python-docx list styles would renumber
                # across the document, and a counter would restart after an interposed formula.
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.9)
                paragraph.paragraph_format.first_line_indent = Cm(-0.5)
                inline(paragraph, f"{bullet.group(2)} {bullet.group(3)}")
            i += 1
            continue
        text = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|```|\$\$|\s*(-|\d+\.)\s)", lines[i].strip()):
            text.append(lines[i].strip())
            i += 1
        inline(document.add_paragraph(), "".join(text))
    document.save(target)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
    print("written", sys.argv[2])
